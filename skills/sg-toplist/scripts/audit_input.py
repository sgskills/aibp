#!/usr/bin/env python3
"""安全读取并审计商品排名输入。

该模块只解析用户明确提供的 XLSX/CSV/TSV，不计算业务结论。普通商品超链接
只计数且永不访问；VBA、外部工作簿连接、公式和单元格提示注入会被拒绝。
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import logging
import math
import re
import sys
import unicodedata
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence
from xml.etree import ElementTree


LOGGER = logging.getLogger(__name__)

SUPPORTED_SUFFIXES = {".xlsx", ".csv", ".tsv"}
MAX_FILE_BYTES = 100 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 512 * 1024 * 1024
MAX_ROWS_PER_SHEET = 200_000
MAX_COLUMNS = 256
MAX_CELLS = 2_000_000
MAX_CELL_CHARS = 100_000
UNDECLARED = "未声明"


class RankingInputError(ValueError):
    """排名输入无法安全、可靠地解析。"""


class UnsafeInputError(RankingInputError):
    """输入包含会改变执行行为的活动内容。"""


class UnsupportedInputError(RankingInputError):
    """输入格式不在白名单内。"""


class DamagedInputError(RankingInputError):
    """文件损坏或结构无法解析。"""


CANONICAL_ALIASES: dict[str, set[str]] = {
    "period": {"日期", "周期", "统计周期", "时间周期", "period", "date", "daterange"},
    "rank": {"行业排名", "商品排名", "市场排名", "排名", "名次", "rank", "ranking"},
    "trend": {"趋势", "平台趋势", "榜单趋势", "trend", "ranktrend", "platformnewflag"},
    "category": {"类目名称", "类目", "品类", "category", "categoryname"},
    "title": {"商品名称", "商品标题", "标题", "productname", "producttitle", "title"},
    "image": {"商品图片", "图片", "image", "productimage"},
    "product_id": {"商品id", "宝贝id", "链接id", "itemid", "productid", "id"},
    "keywords": {"商品关键词", "关键词", "关键字", "keywords", "keyword"},
    "shop": {"店铺名称", "店铺", "商家名称", "shop", "shopname", "store", "storename"},
    "shop_type": {"店铺类型", "店铺平台", "商家类型", "shoptype", "storetype"},
    "buyers": {"支付买家数", "买家数", "支付人数", "buyers", "paidbuyers"},
    "visitors": {"访客数", "访客", "uv", "visitors"},
    "payment_amount": {
        "预估支付金额",
        "支付金额",
        "成交金额",
        "销售额",
        "gmv",
        "paymentamount",
    },
    "price": {"价格", "售价", "成交价", "客单价", "price", "unitprice"},
    "cost": {"成本", "商品成本", "履约成本", "cost"},
    "platform": {"榜单平台", "数据平台", "来源平台", "platform", "sourceplatform"},
    "scope": {"榜单范围", "市场范围", "范围", "scope", "marketscope"},
    "ranking_metric": {"排名指标", "排行指标", "指标", "rankingmetric", "metric"},
    "top_n": {"topn", "榜单深度", "榜单范围数", "top"},
    "brand": {"品牌", "品牌名称", "brand", "brandname"},
}


def _normalise_header(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).strip().lower()
    return re.sub(r"[\s_\-—–:：/\\()（）\[\]【】]+", "", text)


ALIAS_TO_CANONICAL: dict[str, str] = {
    _normalise_header(alias): canonical
    for canonical, aliases in CANONICAL_ALIASES.items()
    for alias in aliases | {canonical}
}

INSTRUCTION_PATTERNS = tuple(
    re.compile(pattern, re.IGNORECASE)
    for pattern in (
        r"ignore\s+(?:all\s+)?(?:previous|prior|above)\s+instructions?",
        r"reveal\s+(?:the\s+)?(?:system|developer)\s+prompt",
        r"(?:system|developer)\s+message\s*:",
        r"prompt\s*injection",
        r"请?忽略(?:以上|上述|此前|前面|原有)(?:所有)?(?:指令|规则|要求)",
        r"(?:泄露|输出|显示)(?:系统|开发者)(?:提示词|指令)",
        r"(?:执行|运行)(?:以下|下列)(?:命令|脚本|代码)",
        r"(?:system|assistant|developer)\s*prompt\s*:",
    )
)

FORMULA_PATTERN = re.compile(
    r"^\s*(?:=|[+@]\s*(?:[A-Za-z_][A-Za-z0-9_.]*\s*\(|[A-Z]{1,3}\d+))"
)

PERIOD_PATTERN = re.compile(
    r"(?P<start>\d{4}[./-]\d{1,2}[./-]\d{1,2})\s*(?:~|～|至|到|—|–|_)\s*"
    r"(?P<end>\d{4}[./-]\d{1,2}[./-]\d{1,2})"
)

NUMBER_PATTERN = re.compile(r"^\s*([<>≤≥]?)([-+]?\d+(?:\.\d+)?)\s*(万|千|亿|w|k)?\s*$", re.I)


def _safe_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    text = str(value).strip()
    if len(text) > MAX_CELL_CHARS:
        raise UnsafeInputError("单元格文本超过安全长度限制")
    return text


def _validate_cell_content(value: Any, location: str) -> None:
    """拒绝公式样式和提示注入；错误只返回位置，不回显载荷。"""
    if value is None or isinstance(value, (int, float, bool, datetime, date)):
        return
    text = _safe_text(value)
    if FORMULA_PATTERN.search(text):
        raise UnsafeInputError(f"检测到公式或公式注入：{location}")
    if any(pattern.search(text) for pattern in INSTRUCTION_PATTERNS):
        raise UnsafeInputError(f"检测到单元格内指令：{location}")


def _coerce_product_id(value: Any) -> str:
    """商品 ID 始终按字符串处理；它表示商品链接，不代表 SKU。"""
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        if value.is_integer():
            return str(int(value))
    return _safe_text(value)


def _coerce_rank(value: Any) -> int | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value if value > 0 else None
    if isinstance(value, float):
        return int(value) if value.is_integer() and value > 0 else None
    text = _safe_text(value)
    if re.fullmatch(r"[1-9]\d*", text):
        return int(text)
    if re.fullmatch(r"[1-9]\d{0,2}(?:,\d{3})+", text):
        return int(text.replace(",", ""))
    # 严禁从 -1、3.5、abc2 等混合文本中截取局部数字。
    return None


def _parse_scaled_number(value: str) -> tuple[float | None, str]:
    match = NUMBER_PATTERN.match(value.replace(",", ""))
    if not match:
        return None, ""
    comparator, number_text, unit = match.groups()
    multiplier = {"万": 10_000, "千": 1_000, "亿": 100_000_000, "w": 10_000, "k": 1_000}.get(
        (unit or "").lower(), 1
    )
    return float(number_text) * multiplier, comparator


def parse_interval(value: Any) -> dict[str, Any] | None:
    """把单值或区间转为上下界；中点明确标记为估算。"""
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        if not math.isfinite(number) or number < 0:
            return None
        return {
            "lower": number,
            "upper": number,
            "estimated_midpoint": number,
            "midpoint_is_estimate": False,
            "raw": value,
        }
    raw = _safe_text(value)
    normalised = raw.replace("～", "~").replace("—", "~").replace("–", "~").replace("至", "~")
    bracket_match = re.match(r"^\s*[\[(（【]\s*(.+?)\s*[,，]\s*(.+?)\s*[\])）】]\s*$", normalised)
    if bracket_match:
        normalised = f"{bracket_match.group(1)}~{bracket_match.group(2)}"
    elif re.match(r"^\s*\d+(?:\.\d+)?\s*(?:万|千|亿|w|k)?\s*-\s*\d", normalised, re.I):
        # 仅在两端都以非负数字开头时把 ASCII 连字符解释为区间，避免误判负数。
        normalised = re.sub(r"\s*-\s*", "~", normalised, count=1)
    parts = [part.strip() for part in normalised.split("~")]
    if len(parts) == 2:
        lower, lower_comparator = _parse_scaled_number(parts[0])
        upper, upper_comparator = _parse_scaled_number(parts[1])
        if lower is None or upper is None or lower_comparator or upper_comparator or lower < 0 or lower > upper:
            return None
        return {
            "lower": lower,
            "upper": upper,
            "estimated_midpoint": (lower + upper) / 2,
            "midpoint_is_estimate": True,
            "raw": raw,
        }
    number, comparator = _parse_scaled_number(raw)
    if number is None or number < 0:
        return None
    if comparator in {"<", "≤"}:
        return {
            "lower": 0.0,
            "upper": number,
            "estimated_midpoint": number / 2,
            "midpoint_is_estimate": True,
            "raw": raw,
        }
    if comparator in {">", "≥"}:
        return {
            "lower": number,
            "upper": None,
            "estimated_midpoint": None,
            "midpoint_is_estimate": True,
            "raw": raw,
        }
    return {
        "lower": number,
        "upper": number,
        "estimated_midpoint": number,
        "midpoint_is_estimate": False,
        "raw": raw,
    }


def interval_change(before: Mapping[str, Any] | None, after: Mapping[str, Any] | None) -> dict[str, Any] | None:
    """按区间边界计算变化范围，不把中点伪装成精确值。"""
    if not before or not after:
        return None
    before_lower = before.get("lower")
    before_upper = before.get("upper")
    after_lower = after.get("lower")
    after_upper = after.get("upper")
    if None in {before_lower, before_upper, after_lower, after_upper}:
        return {"direction": "unknown", "reason": "open_interval"}
    lower_delta = float(after_lower) - float(before_upper)
    upper_delta = float(after_upper) - float(before_lower)
    if lower_delta > 0:
        direction = "increase"
    elif upper_delta < 0:
        direction = "decrease"
    else:
        direction = "overlap_or_uncertain"
    return {
        "lower_delta": lower_delta,
        "upper_delta": upper_delta,
        "direction": direction,
        "ordinal_or_interval_only": True,
    }


def parse_period_label(value: Any) -> dict[str, str | None]:
    raw = _safe_text(value) or UNDECLARED
    match = PERIOD_PATTERN.search(raw)
    if not match:
        return {"label": raw, "start_date": None, "end_date": None}

    def to_iso(text: str) -> str:
        return datetime.strptime(text.replace("/", "-").replace(".", "-"), "%Y-%m-%d").date().isoformat()

    try:
        start_date = to_iso(match.group("start"))
        end_date = to_iso(match.group("end"))
    except ValueError:
        return {"label": raw, "start_date": None, "end_date": None}
    return {"label": f"{start_date} ~ {end_date}", "start_date": start_date, "end_date": end_date}


def _period_sort_key(period: Mapping[str, Any]) -> tuple[str, str, int]:
    if period.get("start_date"):
        return ("0", str(period["start_date"]), int(period.get("input_order") or 0))
    # 未解析周期只保留输入顺序，绝不把标签字符串排序伪装成时间顺序。
    return ("1", "", int(period.get("input_order") or 0))


def _relationship_audit(archive: zipfile.ZipFile) -> dict[str, int]:
    """允许普通商品 hyperlink，但拒绝工作簿外部连接。"""
    hyperlink_count = 0
    for info in archive.infolist():
        name = info.filename.replace("\\", "/").lower()
        if not name.endswith(".rels"):
            continue
        if info.file_size > 10 * 1024 * 1024:
            raise UnsafeInputError("关系文件超过安全长度限制")
        try:
            root = ElementTree.fromstring(archive.read(info))
        except ElementTree.ParseError as exc:
            raise DamagedInputError("XLSX 关系文件损坏") from exc
        for relation in root:
            target_mode = str(relation.attrib.get("TargetMode", "")).lower()
            relation_type = str(relation.attrib.get("Type", "")).lower()
            if target_mode != "external":
                continue
            if relation_type.endswith("/hyperlink"):
                hyperlink_count += 1
                continue
            raise UnsafeInputError("检测到外部工作簿或外部资源连接")
    return {"ordinary_hyperlinks_ignored": hyperlink_count}


def _inspect_xlsx_container(path: Path) -> dict[str, int]:
    try:
        with zipfile.ZipFile(path) as archive:
            corrupt_member = archive.testzip()
            if corrupt_member:
                raise DamagedInputError("XLSX 压缩内容校验失败")
            total_uncompressed = sum(info.file_size for info in archive.infolist())
            if total_uncompressed > MAX_UNCOMPRESSED_BYTES:
                raise UnsafeInputError("XLSX 解压后体积超过安全限制")
            names = {info.filename.replace("\\", "/").lower() for info in archive.infolist()}
            if any("vbaproject" in name or name.endswith(".bin") for name in names):
                raise UnsafeInputError("检测到 VBA 或二进制宏内容")
            if any(
                name.startswith("xl/externallinks/")
                or name == "xl/connections.xml"
                or name.startswith("xl/querytables/")
                for name in names
            ):
                raise UnsafeInputError("检测到外部工作簿连接或查询连接")
            return _relationship_audit(archive)
    except zipfile.BadZipFile as exc:
        raise DamagedInputError("XLSX 文件损坏或并非有效工作簿") from exc


def _read_xlsx(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    security_meta = _inspect_xlsx_container(path)
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise UnsupportedInputError("当前 bundled Python 缺少 openpyxl，无法读取 XLSX；不得静默安装") from exc

    try:
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        raise DamagedInputError("XLSX 工作簿无法解析") from exc

    tables: list[dict[str, Any]] = []
    cell_count = 0
    try:
        for worksheet in workbook.worksheets:
            if worksheet.max_row > MAX_ROWS_PER_SHEET or worksheet.max_column > MAX_COLUMNS:
                raise UnsafeInputError("工作表行列数超过安全限制")
            matrix: list[list[Any]] = []
            for row_number, cells in enumerate(worksheet.iter_rows(), start=1):
                values: list[Any] = []
                for column_number, cell in enumerate(cells, start=1):
                    cell_count += 1
                    if cell_count > MAX_CELLS:
                        raise UnsafeInputError("工作簿单元格数量超过安全限制")
                    if cell.data_type == "f":
                        raise UnsafeInputError(f"检测到公式：{worksheet.title}!R{row_number}C{column_number}")
                    value = cell.value
                    _validate_cell_content(value, f"{worksheet.title}!R{row_number}C{column_number}")
                    values.append(value)
                matrix.append(values)
            if any(any(value not in (None, "") for value in row) for row in matrix):
                tables.append({"sheet": worksheet.title, "matrix": matrix})
    finally:
        workbook.close()
    if not tables:
        raise DamagedInputError("工作簿没有可分析的数据表")
    return tables, security_meta


def _decode_delimited(path: Path) -> str:
    raw = path.read_bytes()
    if b"\x00" in raw:
        raise DamagedInputError("分隔文本包含 NUL 字节")
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DamagedInputError("CSV/TSV 编码无法识别；请使用 UTF-8 或 GB18030")


def _read_delimited(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    text = _decode_delimited(path)
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    matrix: list[list[Any]] = []
    cell_count = 0
    try:
        for row_number, row in enumerate(reader, start=1):
            if row_number > MAX_ROWS_PER_SHEET:
                raise UnsafeInputError("分隔文本行数超过安全限制")
            if len(row) > MAX_COLUMNS:
                raise UnsafeInputError("分隔文本列数超过安全限制")
            values: list[Any] = []
            for column_number, value in enumerate(row, start=1):
                cell_count += 1
                if cell_count > MAX_CELLS:
                    raise UnsafeInputError("分隔文本单元格数量超过安全限制")
                _validate_cell_content(value, f"{path.name}!R{row_number}C{column_number}")
                values.append(value)
            matrix.append(values)
    except csv.Error as exc:
        raise DamagedInputError("CSV/TSV 结构损坏") from exc
    if not matrix:
        raise DamagedInputError("CSV/TSV 没有数据")
    return [{"sheet": path.stem, "matrix": matrix}], {"ordinary_hyperlinks_ignored": 0}


def _detect_header(matrix: Sequence[Sequence[Any]]) -> tuple[int, dict[int, str], list[str]]:
    best: tuple[int, int, dict[int, str], list[str]] | None = None
    for row_index, row in enumerate(matrix[:30]):
        mapping: dict[int, str] = {}
        duplicates: list[str] = []
        seen: set[str] = set()
        for column_index, value in enumerate(row):
            canonical = ALIAS_TO_CANONICAL.get(_normalise_header(value))
            if not canonical:
                continue
            if canonical in seen:
                duplicates.append(canonical)
                continue
            seen.add(canonical)
            mapping[column_index] = canonical
        score = len(mapping) + (3 if "rank" in seen else 0) + (3 if "product_id" in seen else 0)
        candidate = (score, -row_index, mapping, duplicates)
        if best is None or candidate[:2] > best[:2]:
            best = candidate
    if best is None or best[0] < 5 or "rank" not in best[2].values() or "product_id" not in best[2].values():
        raise DamagedInputError("未找到同时包含排名和商品ID的有效表头")
    return -best[1], best[2], best[3]


def _infer_sheet_metadata(sheet_name: str) -> dict[str, str]:
    metadata: dict[str, str] = {}
    for metric in ("交易总量", "支付买家数", "访客数", "预估支付金额", "销售额"):
        if metric in sheet_name:
            metadata["ranking_metric"] = metric
            break
    for scope in ("全网", "仅天猫", "天猫", "淘宝"):
        if scope in sheet_name:
            metadata["scope"] = scope
            break
    if "仅天猫" in sheet_name:
        metadata["platform"] = "天猫"
    elif metadata.get("scope") == "全网":
        # 这是榜单覆盖口径，不是行级店铺类型；不能把天猫/淘宝行拆成两个榜单。
        metadata["platform"] = "全网"
    return metadata


def _table_to_rows(table: Mapping[str, Any], source_name: str) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    matrix = table["matrix"]
    header_index, header_map, duplicate_headers = _detect_header(matrix)
    rows: list[dict[str, Any]] = []
    for row_index, values in enumerate(matrix[header_index + 1 :], start=header_index + 2):
        row: dict[str, Any] = {}
        for column_index, canonical in header_map.items():
            row[canonical] = values[column_index] if column_index < len(values) else None
        if not any(value not in (None, "") for value in row.values()):
            continue
        row["_source"] = {"file": source_name, "sheet": str(table["sheet"]), "row": row_index}
        rows.append(row)
    return rows, {
        "source_file": source_name,
        "sheet": str(table["sheet"]),
        "duplicate_headers": duplicate_headers,
        **_infer_sheet_metadata(str(table["sheet"])),
    }


def _canonical_key(key: Any) -> str:
    text = _normalise_header(key)
    return ALIAS_TO_CANONICAL.get(text, str(key))


def _canonicalize_row(row: Mapping[str, Any], fallback_source: Mapping[str, Any] | None = None) -> dict[str, Any]:
    canonical: dict[str, Any] = {}
    for key, value in row.items():
        if str(key).startswith("_"):
            canonical[str(key)] = value
            continue
        canonical_key = _canonical_key(key)
        if canonical_key not in canonical or canonical[canonical_key] in (None, ""):
            canonical[canonical_key] = value
    source = canonical.get("_source") or fallback_source or {}
    for key, value in canonical.items():
        if key.startswith("_"):
            continue
        _validate_cell_content(value, f"内联数据 {source}")
    canonical["product_id"] = _coerce_product_id(canonical.get("product_id"))
    canonical["rank_raw"] = canonical.get("rank_raw", canonical.get("rank"))
    canonical["rank"] = _coerce_rank(canonical.get("rank"))
    if canonical.get("trend") is True:
        canonical["trend"] = "新上榜"
    for text_field in (
        "period",
        "trend",
        "category",
        "title",
        "keywords",
        "shop",
        "shop_type",
        "brand",
        "platform",
        "scope",
        "ranking_metric",
    ):
        canonical[text_field] = _safe_text(canonical.get(text_field))
    for metric in ("buyers", "visitors", "payment_amount", "price", "cost"):
        raw_value = canonical.get(metric)
        canonical[f"{metric}_raw"] = raw_value
        canonical[f"{metric}_interval"] = parse_interval(raw_value)
    canonical["_source"] = dict(source)
    return canonical


def _metadata_value(metadata: Mapping[str, Any], canonical: str) -> Any:
    for key, value in metadata.items():
        if _canonical_key(key) == canonical and value not in (None, ""):
            return value
    return None


def _normalise_period_blocks(payload: Mapping[str, Any]) -> list[dict[str, Any]]:
    raw_periods = payload.get("periods")
    if raw_periods is None:
        raw_periods = payload.get("snapshots")
    if raw_periods is None and isinstance(payload.get("rows"), Sequence):
        raw_periods = [{"rows": payload.get("rows"), **{k: v for k, v in payload.items() if k != "rows"}}]
    if not isinstance(raw_periods, Sequence) or isinstance(raw_periods, (str, bytes)):
        raise RankingInputError("payload 必须包含 periods 列表或 rows 列表")

    output: list[dict[str, Any]] = []
    global_order = 0
    for block_index, block in enumerate(raw_periods):
        if not isinstance(block, Mapping):
            raise RankingInputError(f"periods[{block_index}] 必须是对象")
        raw_rows = block.get("rows", [])
        if not isinstance(raw_rows, Sequence) or isinstance(raw_rows, (str, bytes)):
            raise RankingInputError(f"periods[{block_index}].rows 必须是列表")
        nested_metadata = block.get("metadata") if isinstance(block.get("metadata"), Mapping) else {}
        metadata = {**nested_metadata, **{k: v for k, v in block.items() if k not in {"rows", "metadata"}}}
        canonical_rows: list[dict[str, Any]] = []
        for row_index, raw_row in enumerate(raw_rows, start=1):
            if not isinstance(raw_row, Mapping):
                raise RankingInputError(f"periods[{block_index}].rows[{row_index - 1}] 必须是对象")
            fallback_source = {
                "file": _safe_text(_metadata_value(metadata, "source_file")) or "inline",
                "sheet": _safe_text(_metadata_value(metadata, "sheet")) or f"period-{block_index + 1}",
                "row": row_index,
            }
            canonical_rows.append(_canonicalize_row(raw_row, fallback_source))

        inferred_period = _safe_text(_metadata_value(metadata, "period"))
        inferred_platform = _safe_text(_metadata_value(metadata, "platform"))
        inferred_scope = _safe_text(_metadata_value(metadata, "scope"))
        inferred_category = _safe_text(_metadata_value(metadata, "category"))
        inferred_metric = _safe_text(_metadata_value(metadata, "ranking_metric"))
        explicit_top_n = _coerce_rank(_metadata_value(metadata, "top_n"))
        source_file = _safe_text(metadata.get("source_file")) or "inline"
        sheet = _safe_text(metadata.get("sheet")) or f"period-{block_index + 1}"
        duplicate_headers = list(metadata.get("duplicate_headers") or [])

        groups: dict[tuple[str, str, str, str, str, int | None], list[dict[str, Any]]] = defaultdict(list)
        for row in canonical_rows:
            period_label = row.get("period") or inferred_period or UNDECLARED
            scope = row.get("scope") or inferred_scope or UNDECLARED
            platform = row.get("platform") or inferred_platform or ("全网" if scope == "全网" else UNDECLARED)
            category = row.get("category") or inferred_category or UNDECLARED
            metric = row.get("ranking_metric") or inferred_metric or UNDECLARED
            row_top_n = _coerce_rank(row.get("top_n")) or explicit_top_n
            groups[(period_label, platform, scope, category, metric, row_top_n)].append(row)

        if not groups and canonical_rows == []:
            groups[(inferred_period or UNDECLARED, inferred_platform or UNDECLARED, inferred_scope or UNDECLARED,
                    inferred_category or UNDECLARED, inferred_metric or UNDECLARED, explicit_top_n)] = []

        for key, rows in groups.items():
            period_label, platform, scope, category, metric, top_n = key
            period_info = parse_period_label(period_label)
            inferred_top_n = max((row["rank"] or 0 for row in rows), default=0) or None
            resolved_top_n = top_n or inferred_top_n
            global_order += 1
            output.append(
                {
                    "period": period_info["label"],
                    "start_date": period_info["start_date"],
                    "end_date": period_info["end_date"],
                    "platform": platform,
                    "scope": scope,
                    "category": category,
                    "ranking_metric": metric,
                    "top_n": resolved_top_n,
                    "top_n_source": metadata.get("top_n_source") or ("declared" if top_n else "max_rank_inferred"),
                    "source_file": source_file,
                    "sheet": sheet,
                    "duplicate_headers": duplicate_headers,
                    "input_order": global_order,
                    "rows": rows,
                }
            )
    return output


def _series_tuple(period: Mapping[str, Any]) -> tuple[str, str, str, str]:
    return tuple(str(period.get(key) or UNDECLARED) for key in ("platform", "scope", "category", "ranking_metric"))


def series_id(period: Mapping[str, Any]) -> str:
    return " | ".join(_series_tuple(period))


def island_id(period: Mapping[str, Any]) -> str:
    return " | ".join(
        [
            *list(_series_tuple(period)),
            str(period.get("period") or UNDECLARED),
            f"Top-{period.get('top_n') or '?'}",
        ]
    )


def normalize_payload(payload: Mapping[str, Any]) -> dict[str, Any]:
    """标准化内联 payload，并保留每行来源位置。"""
    if not isinstance(payload, Mapping):
        raise RankingInputError("payload 必须是对象")
    periods = _normalise_period_blocks(payload)
    for period in periods:
        period["series_id"] = series_id(period)
        period["island_id"] = island_id(period)
        for row in period["rows"]:
            row["_source"]["island_id"] = period["island_id"]
            row["_source"]["series_id"] = period["series_id"]
    return {
        "source_files": list(payload.get("source_files") or []),
        "request": _safe_text(payload.get("request") or payload.get("question") or payload.get("task")),
        "periods": periods,
        "input_controls": dict(payload.get("input_controls") or {}),
    }


def merge_payloads(payloads: Sequence[Mapping[str, Any]]) -> dict[str, Any]:
    periods: list[dict[str, Any]] = []
    source_files: list[str] = []
    controls: Counter[str] = Counter()
    requests: list[str] = []
    for payload in payloads:
        normalised = normalize_payload(payload)
        periods.extend(normalised["periods"])
        source_files.extend(str(path) for path in normalised.get("source_files", []))
        requests.extend([normalised.get("request", "")])
        controls.update({key: int(value) for key, value in normalised.get("input_controls", {}).items() if isinstance(value, int)})
    return {
        "source_files": list(dict.fromkeys(source_files)),
        "request": "\n".join(text for text in requests if text),
        "periods": periods,
        "input_controls": dict(controls),
    }


def load_input(path: str | Path) -> dict[str, Any]:
    """从白名单文件加载排名数据；不会执行任何活动内容或访问链接。"""
    input_path = Path(path).expanduser().resolve(strict=True)
    if not input_path.is_file():
        raise UnsupportedInputError("输入路径不是文件")
    suffix = input_path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise UnsupportedInputError("只支持 XLSX、CSV、TSV；不支持宏工作簿或旧式 XLS")
    if input_path.stat().st_size > MAX_FILE_BYTES:
        raise UnsafeInputError("输入文件超过 100 MiB 安全限制")
    if suffix == ".xlsx":
        tables, controls = _read_xlsx(input_path)
    else:
        tables, controls = _read_delimited(input_path)

    period_blocks: list[dict[str, Any]] = []
    for table in tables:
        rows, metadata = _table_to_rows(table, input_path.name)
        period_blocks.append({**metadata, "rows": rows})
    payload = normalize_payload(
        {
            "source_files": [str(input_path)],
            "periods": period_blocks,
            "input_controls": controls,
        }
    )
    return payload


def _file_rejection(path: str | Path, exc: Exception) -> dict[str, Any]:
    if isinstance(exc, UnsafeInputError):
        code = "UNSAFE_ACTIVE_CONTENT"
    elif isinstance(exc, DamagedInputError):
        code = "DAMAGED_FILE"
    else:
        code = "INPUT_REJECTED"
    return {"path": str(path), "code": code, "message": str(exc)}


def load_inputs_isolated(paths: Sequence[str | Path]) -> dict[str, Any]:
    """逐文件隔离读取；一个坏文件不得抹掉其他安全文件的数据岛。"""
    valid_payloads: list[Mapping[str, Any]] = []
    rejections: list[dict[str, Any]] = []
    for path in paths:
        try:
            valid_payloads.append(load_input(path))
        except (RankingInputError, OSError) as exc:
            rejections.append(_file_rejection(path, exc))
    payload = merge_payloads(valid_payloads) if valid_payloads else {"source_files": [], "periods": [], "input_controls": {}}
    return {"payload": payload, "file_rejections": rejections}


def _issue(code: str, severity: str, message: str, **evidence: Any) -> dict[str, Any]:
    return {"code": code, "severity": severity, "message": message, "evidence": evidence}


def _period_has_blocking_issue(issues: Sequence[Mapping[str, Any]], period_id: str) -> bool:
    structure_blocking_codes = {
        "EMPTY_ISLAND",
        "MISSING_OR_INVALID_RANK",
        "DUPLICATE_RANK",
        "RANK_EXCEEDS_TOP_N",
        "INCOMPLETE_DECLARED_WINDOW",
        "MISSING_PERIOD",
        "REVERSED_PERIOD",
        "DUPLICATE_HEADERS",
        "DUPLICATE_PERIOD_ISLAND",
    }
    return any(
        issue.get("severity") == "error"
        and issue.get("code") in structure_blocking_codes
        and issue.get("evidence", {}).get("island_id") == period_id
        for issue in issues
    )


def build_comparisons(periods: Sequence[Mapping[str, Any]], issues: Sequence[Mapping[str, Any]] | None = None) -> list[dict[str, Any]]:
    """只在同平台/范围/类目/排名指标内建立相邻期比较。"""
    issue_list = list(issues or [])
    grouped: dict[tuple[str, str, str, str], list[Mapping[str, Any]]] = defaultdict(list)
    for period in periods:
        grouped[_series_tuple(period)].append(period)
    comparisons: list[dict[str, Any]] = []
    for _, group in sorted(grouped.items(), key=lambda item: item[0]):
        ordered = sorted(group, key=_period_sort_key)
        for before, after in zip(ordered, ordered[1:]):
            before_top = int(before.get("top_n") or 0)
            after_top = int(after.get("top_n") or 0)
            candidate_common_k = min(before_top, after_top) if before_top and after_top else None
            before_end = before.get("end_date")
            after_start = after.get("start_date")
            dates_parsed = bool(
                before.get("start_date")
                and before_end
                and after_start
                and after.get("end_date")
            )
            periods_forward = bool(
                dates_parsed
                and str(before["start_date"]) <= str(before_end)
                and str(after_start) <= str(after["end_date"])
            )
            if dates_parsed and periods_forward:
                non_overlapping: bool | None = str(before_end) < str(after_start)
            else:
                non_overlapping = None
            before_days = ((date.fromisoformat(str(before_end)) - date.fromisoformat(str(before["start_date"]))).days + 1) if dates_parsed else None
            after_days = ((date.fromisoformat(str(after["end_date"])) - date.fromisoformat(str(after_start))).days + 1) if dates_parsed else None
            same_period_length = bool(periods_forward and before_days == after_days)
            duplicate_period = str(before.get("period")) == str(after.get("period"))
            missing_dimension = any(value == UNDECLARED for value in _series_tuple(before))
            blocking = _period_has_blocking_issue(issue_list, str(before.get("island_id"))) or _period_has_blocking_issue(
                issue_list, str(after.get("island_id"))
            )
            comparable = (
                bool(candidate_common_k)
                and not duplicate_period
                and not missing_dimension
                and dates_parsed
                and periods_forward
                and non_overlapping is not False
                and not blocking
            )
            common_k = candidate_common_k if comparable else None
            reasons: list[str] = []
            if not candidate_common_k:
                reasons.append("missing_top_n")
            if duplicate_period:
                reasons.append("duplicate_period")
            if missing_dimension:
                reasons.append("missing_comparability_dimension")
            if not dates_parsed:
                reasons.append("unparsed_period_dates")
            elif not periods_forward:
                reasons.append("reversed_period")
            if non_overlapping is False:
                reasons.append("overlapping_periods")
            if blocking:
                reasons.append("blocking_data_issue")
            if before_top != after_top and candidate_common_k:
                reasons.append("different_top_n_compare_common_k_only")
            if dates_parsed and not same_period_length:
                reasons.append("different_period_lengths_ordinal_only")
            comparisons.append(
                {
                    "series_id": series_id(before),
                    "before_island": before.get("island_id"),
                    "after_island": after.get("island_id"),
                    "before_period": before.get("period"),
                    "after_period": after.get("period"),
                    "before_top_n": before_top or None,
                    "after_top_n": after_top or None,
                    "common_k": common_k,
                    "candidate_common_k": candidate_common_k,
                    "non_overlapping": non_overlapping,
                    "same_period_length": same_period_length,
                    "before_duration_days": before_days,
                    "after_duration_days": after_days,
                    "continuous": bool(non_overlapping and (date.fromisoformat(str(after_start)) - date.fromisoformat(str(before_end))).days == 1),
                    "comparable": comparable,
                    "reasons": reasons,
                }
            )
    return comparisons


def audit_dataset(payload: Mapping[str, Any]) -> dict[str, Any]:
    """审计结构、主键、排名与可比性；不回显 expected 或生成业务结论。"""
    normalised = normalize_payload(payload)
    periods = normalised["periods"]
    issues: list[dict[str, Any]] = []
    period_summaries: list[dict[str, Any]] = []
    if not periods:
        issues.append(_issue("NO_PERIODS", "error", "没有可分析周期"))

    seen_period_islands: Counter[tuple[str, str]] = Counter()
    for period in periods:
        period_id = str(period["island_id"])
        rows = list(period.get("rows") or [])
        seen_period_islands[(str(period["series_id"]), str(period["period"]))] += 1
        missing_ids = [row for row in rows if not row.get("product_id")]
        missing_ranks = [row for row in rows if row.get("rank") is None]
        ids = [str(row["product_id"]) for row in rows if row.get("product_id")]
        ranks = [int(row["rank"]) for row in rows if row.get("rank") is not None]
        duplicate_ids = sorted(product_id for product_id, count in Counter(ids).items() if count > 1)
        duplicate_ranks = sorted(rank for rank, count in Counter(ranks).items() if count > 1)
        if not rows:
            issues.append(_issue("EMPTY_ISLAND", "error", "可比数据岛没有数据行", island_id=period_id))
        if missing_ids:
            issues.append(
                _issue(
                    "MISSING_PRODUCT_ID",
                    "error",
                    "存在缺失商品ID；仅排除受影响行，不阻断其他ID轨迹",
                    island_id=period_id,
                    count=len(missing_ids),
                    positions=[row.get("_source", {}) for row in missing_ids],
                )
            )
        if missing_ranks:
            issues.append(
                _issue("MISSING_OR_INVALID_RANK", "error", "存在缺失或非法排名", island_id=period_id, count=len(missing_ranks), positions=[{**row.get("_source", {}), "raw_value": row.get("rank_raw")} for row in missing_ranks])
            )
        if duplicate_ids:
            duplicate_positions = {
                product_id: [row.get("_source", {}) for row in rows if str(row.get("product_id") or "") == product_id]
                for product_id in duplicate_ids
            }
            issues.append(
                _issue(
                    "DUPLICATE_PRODUCT_ID",
                    "error",
                    "同一数据岛存在重复商品ID；排除这些ID而非首条胜出",
                    island_id=period_id,
                    count=len(duplicate_ids),
                    sample=duplicate_ids[:5],
                    positions=duplicate_positions,
                )
            )
        if duplicate_ranks:
            issues.append(
                _issue("DUPLICATE_RANK", "error", "同一数据岛存在重复排名", island_id=period_id, count=len(duplicate_ranks), sample=duplicate_ranks[:5], positions=[{**row.get("_source", {}), "raw_value": row.get("rank_raw")} for row in rows if row.get("rank") in duplicate_ranks])
            )
        top_n = int(period.get("top_n") or 0)
        if ranks and top_n and max(ranks) > top_n:
            issues.append(
                _issue("RANK_EXCEEDS_TOP_N", "error", "排名超过声明的 Top-N", island_id=period_id, max_rank=max(ranks), top_n=top_n)
            )
        unique_rank_sequence = sorted(set(ranks))
        if ranks and unique_rank_sequence != list(range(1, max(ranks) + 1)):
            issues.append(
                _issue("RANK_GAPS", "warning", "排名序列存在缺口；保留证据并降级", island_id=period_id, unique_rank_count=len(set(ranks)), max_rank=max(ranks))
            )
        if period.get("top_n_source") == "declared" and top_n:
            expected_ranks = list(range(1, top_n + 1))
            if unique_rank_sequence != expected_ranks:
                issues.append(
                    _issue(
                        "INCOMPLETE_DECLARED_WINDOW",
                        "error",
                        "声明的Top-N窗口不完整；禁止据此计算完整结构、进入或退出",
                        island_id=period_id,
                        declared_top_n=top_n,
                        valid_unique_rank_count=len(unique_rank_sequence),
                        observed_rank_min=min(unique_rank_sequence) if unique_rank_sequence else None,
                        observed_rank_max=max(unique_rank_sequence) if unique_rank_sequence else None,
                    )
                )
        if period.get("period") == UNDECLARED:
            issues.append(_issue("MISSING_PERIOD", "error", "缺失周期口径", island_id=period_id))
        elif not period.get("start_date") or not period.get("end_date"):
            issues.append(
                _issue("UNPARSED_PERIOD", "warning", "周期无法解析为起止日期；禁止纵向比较", island_id=period_id)
            )
        elif str(period["start_date"]) > str(period["end_date"]):
            issues.append(
                _issue("REVERSED_PERIOD", "error", "周期起始日晚于结束日；禁止纵向比较", island_id=period_id)
            )
        for dimension in ("platform", "scope", "category", "ranking_metric"):
            if period.get(dimension) == UNDECLARED:
                issues.append(
                    _issue("MISSING_COMPARABILITY_DIMENSION", "warning", f"缺失可比口径：{dimension}", island_id=period_id, dimension=dimension)
                )
        if period.get("duplicate_headers"):
            issues.append(
                _issue("DUPLICATE_HEADERS", "error", "表头映射到重复标准字段", island_id=period_id, fields=period["duplicate_headers"])
            )
        invalid_intervals: dict[str, int] = {}
        coverage: dict[str, int] = {}
        for metric in ("buyers", "visitors", "payment_amount", "price", "cost"):
            present = sum(row.get(f"{metric}_raw") not in (None, "") for row in rows)
            parsed = sum(row.get(f"{metric}_interval") is not None for row in rows)
            coverage[metric] = present
            if present > parsed:
                invalid_intervals[metric] = present - parsed
        for metric, count in invalid_intervals.items():
            issues.append(
                _issue("INVALID_INTERVAL", "warning", "区间指标无法解析，相关计算降级", island_id=period_id, field=metric, count=count, positions=[{**row.get("_source", {}), "raw_value": row.get(f"{metric}_raw")} for row in rows if row.get(f"{metric}_raw") not in (None, "") and row.get(f"{metric}_interval") is None])
            )
        period_summaries.append(
            {
                "island_id": period_id,
                "series_id": period["series_id"],
                "period": period["period"],
                "start_date": period.get("start_date"),
                "end_date": period.get("end_date"),
                "platform": period["platform"],
                "scope": period["scope"],
                "category": period["category"],
                "ranking_metric": period["ranking_metric"],
                "top_n": period.get("top_n"),
                "top_n_source": period.get("top_n_source"),
                "row_count": len(rows),
                "valid_rank_count": len(ranks),
                "missing_product_id_count": len(missing_ids),
                "duplicate_product_id_count": len(duplicate_ids),
                "missing_rank_count": len(missing_ranks),
                "duplicate_rank_count": len(duplicate_ranks),
                "unique_product_ids": len(set(ids)),
                "unique_ranks": len(set(ranks)),
                "field_coverage": coverage,
                "source_file": period.get("source_file"),
                "sheet": period.get("sheet"),
            }
        )

    for (series_key, period_label), count in seen_period_islands.items():
        if count > 1:
            for summary in period_summaries:
                if summary["series_id"] == series_key and summary["period"] == period_label:
                    issues.append(_issue("DUPLICATE_PERIOD_ISLAND", "error", "同一口径存在重复周期，禁止自动合并", island_id=summary["island_id"], series_id=series_key, period=period_label, count=count))

    for summary in period_summaries:
        local_issues = [issue for issue in issues if issue["evidence"].get("island_id") == summary["island_id"]]
        for issue in local_issues:
            issue["evidence"].update({"file": summary["source_file"], "sheet": summary["sheet"], "period": summary["period"], "denominator": summary["row_count"]})
        summary["issue_counts"] = {severity: sum(issue["severity"] == severity for issue in local_issues) for severity in ("error", "warning")}
        summary["status"] = "blocking" if _period_has_blocking_issue(issues, summary["island_id"]) else ("degraded" if local_issues else "usable")

    comparisons = build_comparisons(periods, issues)
    if any(comparison["non_overlapping"] is False for comparison in comparisons):
        issues.append(_issue("OVERLAPPING_PERIODS", "error", "存在重叠周期，禁止纵向比较"))
    error_count = sum(issue["severity"] == "error" for issue in issues)
    warning_count = sum(issue["severity"] == "warning" for issue in issues)
    local_exclusion_codes = {"MISSING_PRODUCT_ID", "DUPLICATE_PRODUCT_ID"}
    blocking_error_count = sum(
        issue["severity"] == "error" and issue.get("code") not in local_exclusion_codes for issue in issues
    )
    if blocking_error_count:
        status = "blocking"
        conclusion = f"数据审计未通过：{blocking_error_count} 个结构阻断问题、{warning_count} 个警告；仅可报告可验证的输入事实。"
    elif error_count or warning_count:
        status = "degraded"
        conclusion = (
            f"数据可分析但需降级：0 个结构阻断问题、{error_count} 个局部排除问题、"
            f"{warning_count} 个警告。"
        )
    else:
        status = "usable"
        conclusion = "数据结构与主键审计通过，可在相同口径数据岛内分析；排名仍仅是序数证据。"
    return {
        "valid": blocking_error_count == 0,
        "status": status,
        "conclusion": conclusion,
        "error_count": error_count,
        "blocking_error_count": blocking_error_count,
        "warning_count": warning_count,
        "issues": issues,
        "period_count": len({str(period.get("period")) for period in periods}),
        "islands": period_summaries,
        "comparisons": comparisons,
        "input_controls": normalised.get("input_controls", {}),
    }


def audit_paths(paths: Sequence[str | Path]) -> dict[str, Any]:
    """审计多个文件；损坏/活动内容以结构化 blocking 结果返回。"""
    isolated = load_inputs_isolated(paths)
    payload = isolated["payload"]
    rejections = isolated["file_rejections"]
    if not payload.get("periods"):
        first = rejections[0] if rejections else {"code": "INPUT_REJECTED", "message": "没有可读取文件", "path": None}
        return {
            "valid": False,
            "status": "rejected",
            "conclusion": f"输入已拒绝：{first['message']}",
            "error_count": len(rejections) or 1,
            "warning_count": 0,
            "issues": [
                {
                    "code": item["code"],
                    "severity": "error",
                    "message": item["message"],
                    "evidence": {"path": item["path"]},
                }
                for item in rejections
            ]
            or [{"code": "INPUT_REJECTED", "severity": "error", "message": "没有可读取文件", "evidence": {}}],
            "period_count": 0,
            "islands": [],
            "comparisons": [],
            "input_controls": {},
            "file_rejections": rejections,
        }
    result = audit_dataset(payload)
    if rejections:
        safe_status = result["status"]
        result["status"] = "partial"
        result["safe_island_status"] = safe_status
        result["file_rejections"] = rejections
        result["issues"] = list(result.get("issues", [])) + [
            {
                "code": item["code"],
                "severity": "error",
                "message": item["message"],
                "evidence": {"path": item["path"], "isolated_file_only": True},
            }
            for item in rejections
        ]
        result["error_count"] = int(result.get("error_count") or 0) + len(rejections)
        result["conclusion"] = (
            f"{result['conclusion']} 另有 {len(rejections)} 个文件被隔离拒绝；未抹掉安全数据岛。"
        )
    else:
        result["file_rejections"] = []
    return result


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="安全读取并审计商品排名 XLSX/CSV/TSV")
    parser.add_argument("paths", nargs="+", help="一个或多个 XLSX/CSV/TSV 路径")
    parser.add_argument("--format", choices=("json",), default="json")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = _build_parser().parse_args(argv)
    try:
        result = audit_paths(args.paths)
    except (RankingInputError, OSError) as exc:
        LOGGER.error("输入拒绝：%s", exc)
        sys.stdout.write(json.dumps({"status": "rejected", "error": str(exc)}, ensure_ascii=False) + "\n")
        return 2
    sys.stdout.write(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    return 0 if result["status"] in {"usable", "degraded", "partial"} else 1


if __name__ == "__main__":
    raise SystemExit(main())
